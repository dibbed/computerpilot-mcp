"""Native DOCX inspection and conservative mutation primitives."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile

from docx import Document

from core.config import resolve_path
from core.errors import ToolError
from tools.documents.common import artifact_descriptor, publish_document

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_MAX_DOCX_BYTES = 256 * 1024 * 1024
_MAX_ARCHIVE_ENTRIES = 10_000
_MAX_ARCHIVE_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
_MAX_COMPRESSION_RATIO = 1_000
_REQUIRED_MEMBERS = frozenset({"[content_types].xml", "word/document.xml"})


def _docx_target(path: str | Path, *, must_exist: bool) -> Path:
    target = resolve_path(path)
    if target.suffix.casefold() != ".docx":
        raise ToolError(
            "unsupported_docx_format",
            "DOCX tools support .docx files only; macro-enabled .docm mutation is not supported.",
        )
    if must_exist and not target.is_file():
        raise FileNotFoundError(f"File not found: {target}")
    if target.exists() and not target.is_file():
        raise ToolError("not_a_file", f"Target is not a regular file: {target}")
    return target


def _unsafe_archive_name(name: str) -> bool:
    normalized = name.replace("\\", "/")
    path = PurePosixPath(normalized)
    return (
        normalized.startswith(("/", "\\"))
        or (bool(path.parts) and ":" in path.parts[0])
        or ".." in path.parts
    )


def validate_docx_archive(path: str | Path) -> Path:
    target = _docx_target(path, must_exist=True)
    size = target.stat().st_size
    if size > _MAX_DOCX_BYTES:
        raise ToolError(
            "docx_too_large",
            f"DOCX file is {size} bytes; limit is {_MAX_DOCX_BYTES}.",
        )
    try:
        with ZipFile(target, "r") as archive:
            infos = archive.infolist()
            if len(infos) > _MAX_ARCHIVE_ENTRIES:
                raise ToolError(
                    "unsafe_docx_archive",
                    f"DOCX archive has {len(infos)} entries; limit is {_MAX_ARCHIVE_ENTRIES}.",
                )
            names: set[str] = set()
            total_uncompressed = 0
            for info in infos:
                folded = info.filename.casefold()
                if folded in names:
                    raise ToolError(
                        "unsafe_docx_archive",
                        f"DOCX archive contains duplicate member {info.filename!r}.",
                    )
                names.add(folded)
                if _unsafe_archive_name(info.filename):
                    raise ToolError(
                        "unsafe_docx_archive",
                        f"DOCX archive contains unsafe member path {info.filename!r}.",
                    )
                if info.flag_bits & 0x1:
                    raise ToolError(
                        "unsupported_docx_encryption",
                        "Encrypted DOCX archive members are not supported.",
                    )
                total_uncompressed += int(info.file_size)
                if total_uncompressed > _MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise ToolError(
                        "unsafe_docx_archive",
                        "DOCX archive expands beyond the supported uncompressed-size limit.",
                    )
                if info.file_size >= 1_048_576:
                    ratio = info.file_size / max(info.compress_size, 1)
                    if ratio > _MAX_COMPRESSION_RATIO:
                        raise ToolError(
                            "unsafe_docx_archive",
                            "DOCX archive contains an excessive compression ratio.",
                        )
            missing = _REQUIRED_MEMBERS - names
            if missing:
                raise ToolError(
                    "invalid_docx",
                    f"DOCX archive is missing required members: {', '.join(sorted(missing))}.",
                )
            bad_member = archive.testzip()
            if bad_member is not None:
                raise ToolError(
                    "invalid_docx",
                    f"DOCX archive member failed CRC validation: {bad_member!r}.",
                )
    except ToolError:
        raise
    except (BadZipFile, OSError) as exc:
        raise ToolError("invalid_docx", f"DOCX archive is unreadable: {target}") from exc
    return target


def _load_document(path: str | Path) -> tuple[Path, Any]:
    target = validate_docx_archive(path)
    try:
        return target, Document(str(target))
    except Exception as exc:
        raise ToolError("invalid_docx", f"DOCX document could not be opened: {target}") from exc


def _iso_or_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def _bounded_text(value: str, limit: int) -> tuple[str, bool]:
    if limit < 0:
        limit = 0
    return value[:limit], len(value) > limit


def inspect_document(path: str | Path) -> dict[str, Any]:
    target, document = _load_document(path)
    descriptor = artifact_descriptor(target, DOCX_MEDIA_TYPE)
    props = document.core_properties
    core_properties = {
        "title": props.title or "",
        "subject": props.subject or "",
        "author": props.author or "",
        "keywords": props.keywords or "",
        "comments": props.comments or "",
        "created": _iso_or_value(props.created),
        "modified": _iso_or_value(props.modified),
        "last_modified_by": props.last_modified_by or "",
        "category": props.category or "",
    }
    return {
        "ok": True,
        "path": str(target),
        "bytes": descriptor["bytes"],
        "sha256": descriptor["sha256"],
        "media_type": DOCX_MEDIA_TYPE,
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "section_count": len(document.sections),
        "core_properties": core_properties,
        "unsupported_mutations": [
            "comments",
            "headers_footers",
            "tracked_changes",
            "macros",
            "arbitrary_wordprocessingml",
        ],
    }


def _document_items(document: Any) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for index, paragraph in enumerate(document.paragraphs):
        items.append(
            {
                "kind": "paragraph",
                "index": index,
                "text": paragraph.text,
                "style": paragraph.style.name if paragraph.style is not None else None,
            }
        )
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                items.append(
                    {
                        "kind": "table_cell",
                        "table_index": table_index,
                        "row": row_index,
                        "column": column_index,
                        "text": cell.text,
                    }
                )
    return items


def read_document(
    path: str | Path,
    *,
    offset: int = 0,
    max_items: int = 100,
    max_chars: int = 250_000,
) -> dict[str, Any]:
    if offset < 0:
        raise ToolError("invalid_offset", "offset must be zero or greater.")
    if not 1 <= max_items <= 5_000:
        raise ToolError("invalid_max_items", "max_items must be between 1 and 5000.")
    if not 1 <= max_chars <= 2_000_000:
        raise ToolError("invalid_max_chars", "max_chars must be between 1 and 2000000.")

    target, document = _load_document(path)
    all_items = _document_items(document)
    selected = all_items[offset : offset + max_items]
    bounded: list[dict[str, Any]] = []
    remaining = max_chars
    text_truncated = False
    for item in selected:
        if remaining <= 0:
            text_truncated = True
            break
        text = str(item.get("text") or "")
        clipped, clipped_flag = _bounded_text(text, remaining)
        copied = dict(item)
        copied["text"] = clipped
        copied["text_truncated"] = clipped_flag
        bounded.append(copied)
        remaining -= len(clipped)
        if clipped_flag:
            text_truncated = True
            break

    paragraphs = [item for item in bounded if item["kind"] == "paragraph"]
    table_cells = [item for item in bounded if item["kind"] == "table_cell"]
    count = len(bounded)
    total = len(all_items)
    truncated = text_truncated or offset + count < total
    return {
        "ok": True,
        "path": str(target),
        "offset": offset,
        "count": count,
        "total_items": total,
        "has_more": truncated,
        "truncated": truncated,
        "next_offset": offset + count if truncated else None,
        "chars": max_chars - remaining,
        "items": bounded,
        "paragraphs": paragraphs,
        "table_cells": table_cells,
    }


def find_text(
    path: str | Path,
    *,
    query: str,
    case_sensitive: bool = False,
    max_results: int = 100,
    max_items_scanned: int = 100_000,
) -> dict[str, Any]:
    if not query:
        raise ToolError("empty_query", "query must not be empty.")
    if not 1 <= max_results <= 500:
        raise ToolError("invalid_max_results", "max_results must be between 1 and 500.")
    if not 1 <= max_items_scanned <= 1_000_000:
        raise ToolError("invalid_scan_limit", "max_items_scanned must be between 1 and 1000000.")

    target, document = _load_document(path)
    needle = query if case_sensitive else query.casefold()
    matches: list[dict[str, Any]] = []
    scanned = 0
    truncated = False
    for item in _document_items(document):
        if scanned >= max_items_scanned:
            truncated = True
            break
        scanned += 1
        text = str(item.get("text") or "")
        haystack = text if case_sensitive else text.casefold()
        if needle not in haystack:
            continue
        if len(matches) >= max_results:
            truncated = True
            break
        result = dict(item)
        result["text"] = text[:4_000]
        result["text_truncated"] = len(text) > 4_000
        matches.append(result)
    return {
        "ok": True,
        "path": str(target),
        "query": query,
        "case_sensitive": case_sensitive,
        "items_scanned": scanned,
        "count": len(matches),
        "items": matches,
        "truncated": truncated,
    }


def _run_spans(paragraph: Any) -> list[tuple[Any, int, int]]:
    spans: list[tuple[Any, int, int]] = []
    cursor = 0
    for run in paragraph.runs:
        end = cursor + len(run.text)
        spans.append((run, cursor, end))
        cursor = end
    return spans


def _replace_match(paragraph: Any, start: int, end: int, replacement: str) -> None:
    spans = _run_spans(paragraph)
    if not spans:
        raise ToolError(
            "unsupported_docx_paragraph_structure",
            "Matched paragraph has no replaceable text runs.",
        )

    start_index = next(
        (index for index, (_, span_start, span_end) in enumerate(spans) if span_start <= start < span_end),
        None,
    )
    end_position = max(start, end - 1)
    end_index = next(
        (index for index, (_, span_start, span_end) in enumerate(spans) if span_start <= end_position < span_end),
        None,
    )
    if start_index is None or end_index is None:
        raise ToolError(
            "unsupported_docx_paragraph_structure",
            "Matched text could not be mapped to stable text runs.",
        )

    start_run, start_base, _ = spans[start_index]
    end_run, end_base, _ = spans[end_index]
    start_offset = start - start_base
    end_offset = end - end_base

    if start_index == end_index:
        start_run.text = start_run.text[:start_offset] + replacement + start_run.text[end_offset:]
        return

    signatures = {
        str(run._r.rPr.xml) if run._r.rPr is not None else None
        for run, _, _ in spans[start_index : end_index + 1]
    }
    if len(signatures) != 1:
        raise ToolError(
            "docx_mixed_format_match",
            "Matched text spans runs with different formatting; refusing an ambiguous replacement.",
            hint="Replace text within one consistently formatted run or use a narrower match.",
        )

    prefix = start_run.text[:start_offset]
    suffix = end_run.text[end_offset:]
    start_run.text = prefix + replacement
    for index in range(start_index + 1, end_index):
        spans[index][0].text = ""
    end_run.text = suffix


def _replace_paragraph_text(
    paragraph: Any,
    old: str,
    new: str,
    *,
    replace_all: bool,
) -> int:
    source = paragraph.text
    matches: list[tuple[int, int]] = []
    cursor = 0
    while True:
        index = source.find(old, cursor)
        if index < 0:
            break
        matches.append((index, index + len(old)))
        if not replace_all:
            break
        cursor = index + len(old)
    for start, end in reversed(matches):
        _replace_match(paragraph, start, end, new)
    return len(matches)


def _validate_staged_docx(path: Path) -> None:
    validate_docx_archive(path)
    try:
        Document(str(path))
    except Exception as exc:
        raise ToolError("invalid_docx", "Staged DOCX failed round-trip validation.") from exc


def replace_text(
    path: str | Path,
    *,
    old: str,
    new: str,
    replace_all: bool = True,
    include_tables: bool = True,
    expected_sha256: str | None = None,
    backup: bool = True,
) -> dict[str, Any]:
    if not old:
        raise ToolError("empty_old_text", "old text must not be empty.")
    source = _docx_target(path, must_exist=True)
    state = {"replacements": 0}

    def writer(staged: Path) -> None:
        _, document = _load_document(source)
        count = 0
        stop = False
        for paragraph in document.paragraphs:
            replaced = _replace_paragraph_text(
                paragraph,
                old,
                new,
                replace_all=replace_all,
            )
            count += replaced
            if replaced and not replace_all:
                stop = True
                break
        if include_tables and not stop:
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            replaced = _replace_paragraph_text(
                                paragraph,
                                old,
                                new,
                                replace_all=replace_all,
                            )
                            count += replaced
                            if replaced and not replace_all:
                                stop = True
                                break
                        if stop:
                            break
                    if stop:
                        break
                if stop:
                    break
        if count == 0:
            raise ToolError("docx_text_not_found", "Requested text was not found in supported DOCX content.")
        state["replacements"] = count
        document.save(str(staged))

    result = publish_document(
        source,
        writer=writer,
        validator=_validate_staged_docx,
        media_type=DOCX_MEDIA_TYPE,
        backup=backup,
        expected_sha256=expected_sha256,
    )
    return {**result, "replacements": state["replacements"]}


def insert_paragraph(
    path: str | Path,
    *,
    text: str,
    after_index: int | None = None,
    style: str | None = None,
    expected_sha256: str | None = None,
    backup: bool = True,
) -> dict[str, Any]:
    source = _docx_target(path, must_exist=True)
    state: dict[str, Any] = {}

    def writer(staged: Path) -> None:
        _, document = _load_document(source)
        paragraphs = document.paragraphs
        if after_index is None:
            index = len(paragraphs)
            try:
                document.add_paragraph(text, style=style)
            except KeyError as exc:
                raise ToolError("docx_style_not_found", f"DOCX style not found: {style!r}.") from exc
        else:
            if after_index < 0 or after_index >= len(paragraphs):
                raise ToolError(
                    "paragraph_index_out_of_range",
                    f"Paragraph index {after_index} is outside 0..{max(len(paragraphs) - 1, 0)}.",
                )
            index = after_index + 1
            try:
                if index < len(paragraphs):
                    paragraphs[index].insert_paragraph_before(text, style=style)
                else:
                    document.add_paragraph(text, style=style)
            except KeyError as exc:
                raise ToolError("docx_style_not_found", f"DOCX style not found: {style!r}.") from exc
        state["paragraph_index"] = index
        document.save(str(staged))

    result = publish_document(
        source,
        writer=writer,
        validator=_validate_staged_docx,
        media_type=DOCX_MEDIA_TYPE,
        backup=backup,
        expected_sha256=expected_sha256,
    )
    return {**result, **state}


def replace_table_cell(
    path: str | Path,
    *,
    table_index: int,
    row: int,
    column: int,
    text: str,
    expected_sha256: str | None = None,
    backup: bool = True,
) -> dict[str, Any]:
    source = _docx_target(path, must_exist=True)

    def writer(staged: Path) -> None:
        _, document = _load_document(source)
        if table_index < 0 or table_index >= len(document.tables):
            raise ToolError("table_index_out_of_range", f"Table index {table_index} is out of range.")
        table = document.tables[table_index]
        if row < 0 or row >= len(table.rows):
            raise ToolError("table_row_out_of_range", f"Table row {row} is out of range.")
        if column < 0 or column >= len(table.columns):
            raise ToolError("table_column_out_of_range", f"Table column {column} is out of range.")
        cell = table.cell(row, column)
        if len(cell.paragraphs) != 1:
            raise ToolError(
                "unsupported_docx_cell_structure",
                "Table-cell replacement currently requires exactly one paragraph to preserve formatting safely.",
            )
        paragraph = cell.paragraphs[0]
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(text)
        document.save(str(staged))

    result = publish_document(
        source,
        writer=writer,
        validator=_validate_staged_docx,
        media_type=DOCX_MEDIA_TYPE,
        backup=backup,
        expected_sha256=expected_sha256,
    )
    return {
        **result,
        "table_index": table_index,
        "row": row,
        "column": column,
        "text": text,
    }


def create_document(
    path: str | Path,
    *,
    paragraphs: list[str],
    title: str | None = None,
    overwrite: bool = False,
    expected_sha256: str | None = None,
    backup: bool = True,
    create_parents: bool = False,
) -> dict[str, Any]:
    target = _docx_target(path, must_exist=False)
    if expected_sha256 is not None and not overwrite:
        raise ToolError(
            "expected_hash_requires_overwrite",
            "expected_sha256 is only meaningful when overwrite=True.",
        )

    def writer(staged: Path) -> None:
        document = Document()
        if title is not None:
            document.core_properties.title = title
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        document.save(str(staged))

    result = publish_document(
        target,
        writer=writer,
        validator=_validate_staged_docx,
        media_type=DOCX_MEDIA_TYPE,
        backup=backup,
        expected_sha256=expected_sha256,
        create_parents=create_parents,
        exclusive_create=not overwrite,
    )
    return {**result, "paragraph_count": len(paragraphs)}
