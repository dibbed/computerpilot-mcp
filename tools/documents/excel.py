"""Native Excel inspection primitives backed by openpyxl."""

from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from openpyxl.utils.exceptions import InvalidFileException

from core.config import resolve_path
from core.errors import ToolError
from core.response import page
from tools.documents.common import artifact_descriptor

EXCEL_MEDIA_TYPES = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
}


def excel_path(path: str | Path) -> Path:
    target = resolve_path(path)
    suffix = target.suffix.casefold()
    if suffix not in EXCEL_MEDIA_TYPES:
        raise ToolError(
            "unsupported_excel_format",
            "Excel tools support .xlsx and .xlsm files.",
        )
    if not target.is_file():
        raise FileNotFoundError(f"File not found: {target}")
    return target


def media_type(path: Path) -> str:
    return EXCEL_MEDIA_TYPES[path.suffix.casefold()]


def has_vba_project(path: Path) -> bool:
    try:
        with ZipFile(path, "r") as archive:
            names = {name.casefold() for name in archive.namelist()}
    except (BadZipFile, OSError) as exc:
        raise ToolError("invalid_excel", f"Workbook archive is unreadable: {path}") from exc
    return "xl/vbaproject.bin" in names


def load_excel(path: str | Path, *, data_only: bool = False) -> tuple[Path, Any]:
    target = excel_path(path)
    try:
        workbook = load_workbook(
            filename=target,
            read_only=False,
            data_only=data_only,
            keep_vba=target.suffix.casefold() == ".xlsm",
        )
    except (InvalidFileException, BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ToolError("invalid_excel", f"Workbook could not be opened: {target}") from exc
    return target, workbook


def _sheet_summary(sheet: Any, index: int) -> dict[str, Any]:
    return {
        "index": index,
        "name": sheet.title,
        "state": sheet.sheet_state,
        "max_row": int(sheet.max_row),
        "max_column": int(sheet.max_column),
        "merged_range_count": len(sheet.merged_cells.ranges),
        "table_count": len(sheet.tables),
    }


def inspect_workbook(path: str | Path, *, max_sheets: int = 100) -> dict[str, Any]:
    if not 1 <= max_sheets <= 500:
        raise ToolError("invalid_max_sheets", "max_sheets must be between 1 and 500.")
    target, workbook = load_excel(path)
    try:
        total = len(workbook.worksheets)
        sheets = [
            _sheet_summary(sheet, index)
            for index, sheet in enumerate(workbook.worksheets[:max_sheets])
        ]
        descriptor = artifact_descriptor(target, media_type(target))
        return {
            "ok": True,
            "path": str(target),
            "format": target.suffix.casefold().lstrip("."),
            "bytes": descriptor["bytes"],
            "sha256": descriptor["sha256"],
            "media_type": descriptor["media_type"],
            "sheet_count": total,
            "sheets": sheets,
            "truncated": len(sheets) < total,
            "has_macros": has_vba_project(target),
        }
    finally:
        workbook.close()


def list_sheets(
    path: str | Path,
    *,
    offset: int = 0,
    max_items: int = 50,
) -> dict[str, Any]:
    if offset < 0:
        raise ToolError("invalid_offset", "offset must be zero or greater.")
    if not 1 <= max_items <= 500:
        raise ToolError("invalid_max_items", "max_items must be between 1 and 500.")
    target, workbook = load_excel(path)
    try:
        rows = [
            _sheet_summary(sheet, index)
            for index, sheet in enumerate(workbook.worksheets)
        ]
        total = len(rows)
        return {
            "ok": True,
            "path": str(target),
            **page(
                rows[offset : offset + max_items],
                total=total,
                offset=offset,
                limit=max_items,
            ),
        }
    finally:
        workbook.close()


def _worksheet(workbook: Any, sheet: str) -> Any:
    if sheet not in workbook.sheetnames:
        raise ToolError("sheet_not_found", f"Worksheet not found: {sheet!r}.")
    return workbook[sheet]


def _json_scalar(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _bounded_scalar(value: Any, remaining_chars: int) -> tuple[Any, int, bool]:
    normalized = _json_scalar(value)
    if normalized is None:
        return None, 0, False
    rendered = normalized if isinstance(normalized, str) else str(normalized)
    if isinstance(normalized, str) and len(rendered) > max(remaining_chars, 0):
        allowed = max(remaining_chars, 0)
        return normalized[:allowed], allowed, True
    return normalized, min(len(rendered), max(remaining_chars, 0)), False


def read_range(
    path: str | Path,
    *,
    sheet: str,
    cell_range: str,
    max_cells: int = 5_000,
    max_chars: int = 250_000,
    data_only: bool = False,
) -> dict[str, Any]:
    if not 1 <= max_cells <= 50_000:
        raise ToolError("invalid_max_cells", "max_cells must be between 1 and 50000.")
    if not 1 <= max_chars <= 2_000_000:
        raise ToolError("invalid_max_chars", "max_chars must be between 1 and 2000000.")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except (TypeError, ValueError) as exc:
        raise ToolError("invalid_cell_range", f"Invalid Excel range: {cell_range!r}.") from exc
    row_count = max_row - min_row + 1
    column_count = max_col - min_col + 1
    cell_count = row_count * column_count
    if cell_count > max_cells:
        raise ToolError(
            "range_too_large",
            f"Requested range contains {cell_count} cells; max_cells is {max_cells}.",
            hint="Read a smaller range or raise max_cells within the tool limit.",
        )

    target, workbook = load_excel(path, data_only=data_only)
    try:
        worksheet = _worksheet(workbook, sheet)
        values: list[list[Any]] = []
        remaining = max_chars
        truncated_cells = 0
        for row in worksheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        ):
            output_row: list[Any] = []
            for cell in row:
                normalized, used, truncated = _bounded_scalar(cell.value, remaining)
                remaining = max(0, remaining - used)
                truncated_cells += int(truncated)
                output_row.append(normalized)
            values.append(output_row)
        return {
            "ok": True,
            "path": str(target),
            "sheet": sheet,
            "range": cell_range,
            "row_count": row_count,
            "column_count": column_count,
            "cell_count": cell_count,
            "values": values,
            "data_only": data_only,
            "truncated": truncated_cells > 0,
            "truncated_cells": truncated_cells,
            "max_chars": max_chars,
        }
    finally:
        workbook.close()


def find_cells(
    path: str | Path,
    *,
    query: str,
    sheet: str | None = None,
    case_sensitive: bool = False,
    max_results: int = 100,
    max_cells_scanned: int = 100_000,
) -> dict[str, Any]:
    if not query:
        raise ToolError("empty_query", "query must not be empty.")
    if not 1 <= max_results <= 500:
        raise ToolError("invalid_max_results", "max_results must be between 1 and 500.")
    if not 1 <= max_cells_scanned <= 1_000_000:
        raise ToolError(
            "invalid_scan_limit",
            "max_cells_scanned must be between 1 and 1000000.",
        )

    target, workbook = load_excel(path)
    try:
        worksheets = [_worksheet(workbook, sheet)] if sheet is not None else list(workbook.worksheets)
        needle = query if case_sensitive else query.casefold()
        items: list[dict[str, Any]] = []
        scanned = 0
        truncated = False
        stop = False
        for worksheet in worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if scanned >= max_cells_scanned:
                        truncated = True
                        stop = True
                        break
                    scanned += 1
                    if cell.value is None:
                        continue
                    rendered = str(_json_scalar(cell.value))
                    haystack = rendered if case_sensitive else rendered.casefold()
                    if needle not in haystack:
                        continue
                    if len(items) >= max_results:
                        truncated = True
                        stop = True
                        break
                    bounded = rendered[:4_000]
                    items.append({
                        "sheet": worksheet.title,
                        "coordinate": cell.coordinate,
                        "value": bounded,
                        "value_truncated": len(bounded) < len(rendered),
                        "is_formula": isinstance(cell.value, str) and cell.value.startswith("="),
                    })
                if stop:
                    break
            if stop:
                break
        return {
            "ok": True,
            "path": str(target),
            "query": query,
            "case_sensitive": case_sensitive,
            "cells_scanned": scanned,
            "count": len(items),
            "items": items,
            "truncated": truncated,
        }
    finally:
        workbook.close()


def get_formula(path: str | Path, *, sheet: str, cell: str) -> dict[str, Any]:
    try:
        coordinate_to_tuple(cell)
    except (TypeError, ValueError) as exc:
        raise ToolError("invalid_cell", f"Invalid Excel cell coordinate: {cell!r}.") from exc
    target, workbook = load_excel(path)
    try:
        worksheet = _worksheet(workbook, sheet)
        resolved = worksheet[cell]
        value = resolved.value
        is_formula = isinstance(value, str) and value.startswith("=")
        return {
            "ok": True,
            "path": str(target),
            "sheet": sheet,
            "cell": resolved.coordinate,
            "is_formula": is_formula,
            "formula": value if is_formula else None,
        }
    finally:
        workbook.close()


def table_info(
    path: str | Path,
    *,
    sheet: str | None = None,
    offset: int = 0,
    max_items: int = 50,
) -> dict[str, Any]:
    if offset < 0:
        raise ToolError("invalid_offset", "offset must be zero or greater.")
    if not 1 <= max_items <= 500:
        raise ToolError("invalid_max_items", "max_items must be between 1 and 500.")
    target, workbook = load_excel(path)
    try:
        worksheets = [_worksheet(workbook, sheet)] if sheet is not None else list(workbook.worksheets)
        tables: list[dict[str, Any]] = []
        for worksheet in worksheets:
            for table in worksheet.tables.values():
                style = table.tableStyleInfo
                tables.append({
                    "sheet": worksheet.title,
                    "name": table.displayName,
                    "range": table.ref,
                    "style": None if style is None else style.name,
                })
        total = len(tables)
        return {
            "ok": True,
            "path": str(target),
            **page(
                tables[offset : offset + max_items],
                total=total,
                offset=offset,
                limit=max_items,
            ),
        }
    finally:
        workbook.close()
