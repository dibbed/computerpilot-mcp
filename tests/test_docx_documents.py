from __future__ import annotations

import asyncio
import importlib
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from core.errors import ToolError
from core.registry import create_server
from core.tooling import MUTATING_TOOL_OPERATIONS


def _document(path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("Hello ")
    first.bold = True
    second = paragraph.add_run("world")
    second.italic = True
    document.add_paragraph("Second paragraph", style="Quote")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    cell_run = table.cell(1, 0).paragraphs[0].add_run("Alpha")
    cell_run.bold = True
    table.cell(1, 1).text = "42"
    document.core_properties.title = "Sample"
    document.save(str(path))


def test_docx_inspect_read_and_find_are_bounded_and_structured(tmp_path: Path) -> None:
    docx_tools = importlib.import_module("tools.documents.docx")
    path = tmp_path / "sample.docx"
    _document(path)

    inspected = docx_tools.inspect_document(path)
    assert inspected["paragraph_count"] == 2
    assert inspected["table_count"] == 1
    assert inspected["section_count"] == 1
    assert inspected["core_properties"]["title"] == "Sample"
    assert inspected["media_type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    read = docx_tools.read_document(path, offset=0, max_items=10, max_chars=25)
    assert read["paragraphs"][0]["text"] == "Hello world"
    assert read["truncated"] is True
    assert read["chars"] <= 25
    assert read["total_items"] >= 4

    found = docx_tools.find_text(path, query="alpha", max_results=10)
    assert found["count"] == 1
    assert found["items"][0]["kind"] == "table_cell"
    assert found["items"][0]["table_index"] == 0
    assert found["items"][0]["row"] == 1
    assert found["items"][0]["column"] == 0


def test_docx_replacement_preserves_supported_run_formatting_and_is_atomic(tmp_path: Path) -> None:
    docx_tools = importlib.import_module("tools.documents.docx")
    path = tmp_path / "replace.docx"
    _document(path)

    result = docx_tools.replace_text(
        path,
        old="Hello",
        new="Greetings",
        replace_all=False,
        backup=True,
    )
    assert result["replacements"] == 1
    assert Path(result["backup"]).is_file()
    assert result["artifact"]["path"] == str(path)
    assert result["postcondition"]["kind"] == "file_sha256"

    document = Document(str(path))
    try:
        paragraph = document.paragraphs[0]
        assert paragraph.text == "Greetings world"
        assert paragraph.runs[0].bold is True
        assert paragraph.runs[-1].italic is True
    finally:
        del document


def test_docx_insert_table_replace_and_create_round_trip(tmp_path: Path) -> None:
    docx_tools = importlib.import_module("tools.documents.docx")
    path = tmp_path / "mutate.docx"
    _document(path)

    inserted = docx_tools.insert_paragraph(
        path,
        text="Inserted",
        after_index=0,
        style="Quote",
        backup=False,
    )
    assert inserted["paragraph_index"] == 1

    replaced = docx_tools.replace_table_cell(
        path,
        table_index=0,
        row=1,
        column=0,
        text="Beta",
        expected_sha256=inserted["sha256"],
        backup=False,
    )
    assert replaced["text"] == "Beta"

    document = Document(str(path))
    try:
        assert document.paragraphs[1].text == "Inserted"
        assert document.paragraphs[1].style is not None
        assert document.paragraphs[1].style.name == "Quote"
        assert document.tables[0].cell(1, 0).text == "Beta"
        assert document.tables[0].cell(1, 0).paragraphs[0].runs[0].bold is True
    finally:
        del document

    created_path = tmp_path / "created.docx"
    created = docx_tools.create_document(
        created_path,
        paragraphs=["One", "Two"],
        title="Created",
        create_parents=False,
    )
    assert created["created"] is True
    assert created["artifact"]["media_type"].endswith("wordprocessingml.document")
    created_doc = Document(str(created_path))
    try:
        assert [p.text for p in created_doc.paragraphs] == ["One", "Two"]
        assert created_doc.core_properties.title == "Created"
    finally:
        del created_doc


def test_docx_rejects_malformed_or_unsafe_archives(tmp_path: Path) -> None:
    docx_tools = importlib.import_module("tools.documents.docx")

    malformed = tmp_path / "malformed.docx"
    with ZipFile(malformed, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
    with pytest.raises(ToolError) as error:
        docx_tools.inspect_document(malformed)
    assert error.value.code == "invalid_docx"

    traversal = tmp_path / "traversal.docx"
    with ZipFile(traversal, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("../escape.bin", b"x")
    with pytest.raises(ToolError) as unsafe:
        docx_tools.inspect_document(traversal)
    assert unsafe.value.code == "unsafe_docx_archive"


def test_docx_replace_fails_closed_on_mixed_format_cross_run_match(tmp_path: Path) -> None:
    docx_tools = importlib.import_module("tools.documents.docx")
    path = tmp_path / "mixed.docx"
    _document(path)
    before = path.read_bytes()

    with pytest.raises(ToolError) as error:
        docx_tools.replace_text(
            path,
            old="Hello world",
            new="Greetings earth",
            replace_all=False,
            backup=False,
        )

    assert error.value.code == "docx_mixed_format_match"
    assert path.read_bytes() == before


def test_docx_tools_are_registered_and_mutations_are_recovery_guarded() -> None:
    names = {tool.name for tool in asyncio.run(create_server().list_tools())}
    assert {
        "docx_inspect",
        "docx_read",
        "docx_find",
        "docx_replace_text",
        "docx_insert_paragraph",
        "docx_replace_table_cell",
        "docx_create",
    } <= names
    assert {
        "docx_replace_text",
        "docx_insert_paragraph",
        "docx_replace_table_cell",
        "docx_create",
    } <= MUTATING_TOOL_OPERATIONS


def test_docx_replace_fails_closed_when_text_is_missing(tmp_path: Path) -> None:
    docx_tools = importlib.import_module("tools.documents.docx")
    path = tmp_path / "missing.docx"
    _document(path)
    before = path.read_bytes()

    with pytest.raises(ToolError) as error:
        docx_tools.replace_text(
            path,
            old="does not exist",
            new="replacement",
            backup=False,
        )
    assert error.value.code == "docx_text_not_found"
    assert path.read_bytes() == before
