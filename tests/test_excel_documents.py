from __future__ import annotations

import asyncio
import importlib
import importlib.util
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.registry import create_server


def _workbook(path: Path) -> None:
    workbook = Workbook()
    active = workbook.active
    active.title = "Summary"
    active["A1"] = "Name"
    active["B1"] = "Value"
    active["A2"] = "Alpha"
    active["B2"] = "=1+1"
    active.merge_cells("D1:E1")
    active["D1"] = "Merged"
    data = workbook.create_sheet("Data")
    data.sheet_state = "hidden"
    data["A1"] = 42
    workbook.save(path)
    workbook.close()


def test_excel_inspect_and_list_sheets_use_real_workbook_metadata(tmp_path: Path) -> None:
    assert importlib.util.find_spec("tools.documents.excel") is not None
    excel = importlib.import_module("tools.documents.excel")
    inspect_workbook = getattr(excel, "inspect_workbook", None)
    list_sheets = getattr(excel, "list_sheets", None)
    assert callable(inspect_workbook)
    assert callable(list_sheets)

    path = tmp_path / "sample.xlsx"
    _workbook(path)

    inspected = inspect_workbook(path, max_sheets=10)
    assert inspected["ok"] is True
    assert inspected["format"] == "xlsx"
    assert inspected["sheet_count"] == 2
    assert inspected["truncated"] is False
    assert inspected["has_macros"] is False
    assert inspected["sheets"][0] == {
        "index": 0,
        "name": "Summary",
        "state": "visible",
        "max_row": 2,
        "max_column": 5,
        "merged_range_count": 1,
        "table_count": 0,
    }
    assert inspected["sheets"][1]["name"] == "Data"
    assert inspected["sheets"][1]["state"] == "hidden"

    listed = list_sheets(path, offset=1, max_items=1)
    assert listed["total_count"] == 2
    assert listed["count"] == 1
    assert listed["items"][0]["name"] == "Data"
    assert listed["has_more"] is False


def test_excel_inspect_and_list_sheets_are_registered_tools() -> None:
    tools = asyncio.run(create_server().list_tools())
    names = {tool.name for tool in tools}
    assert {"excel_inspect", "excel_list_sheets"} <= names


def test_excel_bounded_reads_find_formula_and_table_info(tmp_path: Path) -> None:
    excel = importlib.import_module("tools.documents.excel")
    read_range = getattr(excel, "read_range", None)
    find_cells = getattr(excel, "find_cells", None)
    get_formula = getattr(excel, "get_formula", None)
    table_info = getattr(excel, "table_info", None)
    assert callable(read_range)
    assert callable(find_cells)
    assert callable(get_formula)
    assert callable(table_info)

    path = tmp_path / "reads.xlsx"
    _workbook(path)
    workbook = load_workbook(path)
    sheet = workbook["Summary"]
    table = Table(displayName="SummaryTable", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet["C3"] = "x" * 30
    workbook.save(path)
    workbook.close()

    ranged = read_range(
        path,
        sheet="Summary",
        cell_range="A1:C3",
        max_cells=20,
        max_chars=20,
    )
    assert ranged["row_count"] == 3
    assert ranged["column_count"] == 3
    assert ranged["values"][0] == ["Name", "Value", None]
    assert ranged["values"][1][1] == "=1+1"
    assert ranged["truncated"] is True
    assert ranged["truncated_cells"] == 1

    found = find_cells(
        path,
        query="alpha",
        sheet="Summary",
        case_sensitive=False,
        max_results=10,
        max_cells_scanned=100,
    )
    assert found["count"] == 1
    assert found["items"][0]["coordinate"] == "A2"
    assert found["items"][0]["value"] == "Alpha"
    assert found["truncated"] is False

    formula = get_formula(path, sheet="Summary", cell="B2")
    assert formula["is_formula"] is True
    assert formula["formula"] == "=1+1"

    tables = table_info(path, sheet="Summary", offset=0, max_items=10)
    assert tables["total_count"] == 1
    assert tables["items"][0]["name"] == "SummaryTable"
    assert tables["items"][0]["range"] == "A1:B2"


def test_excel_read_tools_are_registered() -> None:
    names = {tool.name for tool in asyncio.run(create_server().list_tools())}
    assert {
        "excel_read_range",
        "excel_find",
        "excel_get_formula",
        "excel_table_info",
    } <= names
